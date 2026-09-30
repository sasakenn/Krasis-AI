"""
task_generator.py
------------------
「task-generator」チャンネル用: 課題の説明文から、実際に提出できる成果物を生成する。

task_estimator.py(所要時間の見積もりだけ)とは異なり、こちらは成果物そのものを作る。
- generate_task_text: レポート・回答など、文章で完結する課題の内容を生成する。
- generate_task_spreadsheet: Excel(表計算)が必要な課題について、表の構造をJSONで
  Claudeに設計させる。実際の.xlsxバイト列への変換は build_xlsx_bytes が行う
  (Claudeはバイナリを直接出力できないため、JSON仕様 → openpyxlで組み立てる二段構え)。
"""

from __future__ import annotations

import io
import json
import logging
import os
import re

from anthropic import Anthropic
from dotenv import load_dotenv
from openpyxl import Workbook

from lang_utils import fallback_answer, normalize_lang

load_dotenv()

logger = logging.getLogger(__name__)

ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY")
MODEL = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-5")
client = Anthropic(api_key=ANTHROPIC_API_KEY) if ANTHROPIC_API_KEY else None

# 後方互換用(既存のimport元向け)。実際の表示にはlang別のfallback_answer()を使う。
FALLBACK_TEXT = (
    "現在この機能を利用できません(APIキー未設定、または一時的なエラーです)。"
    "しばらく待ってから再試行してください。"
)

_TASK_TEXT_REQUIREMENTS = {
    "ja": "- 日本語で執筆する",
    "en": "- Write in English",
    "ko": "- 한국어로 작성할 것",
}

_INVALID_SHEET_NAME_CHARS = re.compile(r"[\\/*?:\[\]]")

# openpyxlは"="で始まる文字列セルを実際のExcel数式として書き込む(+-@始まりは通常の
# 文字列のままになることを確認済み。CSVでの挙動とは異なる)。Claudeへの指示上、
# 集計用にSUM等の数式文字列を使わせているため、その安全な形だけを許可し、それ以外の
# "="始まり文字列は数式インジェクション対策として先頭に"'"を付けて無害化する
# (description=課題文はユーザー入力なので、Claudeの応答を経由した間接的な攻撃を想定)。
_SAFE_FORMULA_PATTERN = re.compile(
    r"^=(SUM|AVERAGE|COUNT|COUNTA|MIN|MAX)\(\$?[A-Za-z]{1,3}\$?[0-9]+(:\$?[A-Za-z]{1,3}\$?[0-9]+)?\)$"
)


def _sanitize_cell_value(value):
    if isinstance(value, str) and value.startswith("=") and not _SAFE_FORMULA_PATTERN.match(value.strip()):
        return "'" + value
    return value


def _parse_json_response(raw_text: str) -> dict:
    text = raw_text.strip()
    if text.startswith("```"):
        text = text.split("```", 2)[1].replace("json", "", 1).strip() if "```" in text else text
    return json.loads(text)


def generate_task_text(description: str, lang: str = "ja") -> dict:
    """文章で完結する課題(レポート・回答・要約など)の内容を生成する。langで指定した
    言語(ja/en/ko)で執筆させる。

    Claudeが未設定/呼び出し失敗の場合は、その旨を伝えるフォールバック回答を返す
    (course_guide.answer_course_questionと同じ方針)。
    """
    no_usage = {"input_tokens": 0, "output_tokens": 0}
    clean = description.strip()
    lang = normalize_lang(lang)

    if client is None:
        return {"content": fallback_answer(lang), "_token_usage": no_usage}

    language_requirement = _TASK_TEXT_REQUIREMENTS[lang]
    prompt = f"""あなたは大学生の課題作成を支援するアシスタントです。

以下の課題に、実際に提出できる完成した文章で取り組んでください(構成案ではなく完成した内容にすること)。

課題内容: {clean}

要件:
{language_requirement}
- Markdown形式で、必要に応じて見出し(## )や箇条書きを使う
- 前後に説明文やコードブロック記号は付けず、本文のみを返す"""

    try:
        resp = client.messages.create(
            model=MODEL,
            max_tokens=4000,
            messages=[{"role": "user", "content": prompt}],
        )
        text_block = next((b for b in resp.content if b.type == "text"), None)
        content = text_block.text.strip() if text_block else fallback_answer(lang)
        return {
            "content": content,
            "_token_usage": {
                "input_tokens": resp.usage.input_tokens,
                "output_tokens": resp.usage.output_tokens,
            },
        }
    except Exception:
        logger.warning(
            "Claude呼び出しに失敗したため課題生成のフォールバック回答を返します(description=%r)",
            clean,
            exc_info=True,
        )
        return {"content": fallback_answer(lang), "_token_usage": no_usage}


def _build_spreadsheet_prompt(description: str, lang: str) -> str:
    language_requirement = _TASK_TEXT_REQUIREMENTS[normalize_lang(lang)]
    return f"""あなたは大学生の課題作成を支援するアシスタントで、Excel(表計算)の課題を担当します。

以下の課題を満たす表を設計してください。

課題内容: {description}

要件:
- JSONのみを返す(前後に説明文やコードブロック記号は付けない)
{language_requirement}(見出し・セルの文字列に使う言語)
- 出力形式:
{{
  "filename": "ファイル名(拡張子なし、日本語可)",
  "sheets": [
    {{
      "name": "シート名(31文字以内)",
      "headers": ["列見出し1", "列見出し2"],
      "rows": [
        ["セルの値", "セルの値"]
      ]
    }}
  ]
}}
- セルの値は、数値はJSONの数値型、文字列はJSON文字列型にする
- 合計・平均などが必要な場合は、セルの値をExcelの数式文字列(例: "=SUM(B2:B13)")にしてよい
  (見出し行が1行目になるので、データはheadersの次の行=2行目から始まる前提でセル位置を数えること)
- 課題に応じて複数シートにしてよい(例: 元データと集計を分ける)
- 架空のデータで構わないが、課題の要件(項目数・粒度など)を満たす具体的な内容にする"""


_FALLBACK_SPREADSHEET_TEXT = {
    "ja": {"filename": "課題", "header": "項目", "cell": "現在この機能を利用できません"},
    "en": {"filename": "assignment", "header": "Item", "cell": "This feature is currently unavailable"},
    "ko": {"filename": "과제", "header": "항목", "cell": "현재 이 기능을 사용할 수 없습니다"},
}


def _fallback_spreadsheet_spec(lang: str = "ja") -> dict:
    text = _FALLBACK_SPREADSHEET_TEXT[normalize_lang(lang)]
    return {
        "filename": text["filename"],
        "sheets": [
            {
                "name": "Sheet1",
                "headers": [text["header"]],
                "rows": [[text["cell"]]],
            }
        ],
    }


def generate_task_spreadsheet(description: str, lang: str = "ja") -> dict:
    """Excel(表計算)が必要な課題について、表の構造(spec)をJSONで生成する。langで
    指定した言語(ja/en/ko)で見出し・セルの文字列を書かせる。

    戻り値の "spec" を build_xlsx_bytes に渡すと、実際の.xlsxバイト列になる。
    Claudeが未設定/呼び出し失敗/JSON解析失敗の場合は、フォールバックの spec を返す。
    """
    no_usage = {"input_tokens": 0, "output_tokens": 0}
    clean = description.strip()
    lang = normalize_lang(lang)

    if client is None:
        return {"spec": _fallback_spreadsheet_spec(lang), "_token_usage": no_usage}

    try:
        resp = client.messages.create(
            model=MODEL,
            max_tokens=4000,
            messages=[{"role": "user", "content": _build_spreadsheet_prompt(clean, lang)}],
        )
        text_block = next((b for b in resp.content if b.type == "text"), None)
        spec = _parse_json_response(text_block.text) if text_block else _fallback_spreadsheet_spec(lang)
        return {
            "spec": spec,
            "_token_usage": {
                "input_tokens": resp.usage.input_tokens,
                "output_tokens": resp.usage.output_tokens,
            },
        }
    except Exception:
        logger.warning(
            "Claude呼び出しに失敗したため表計算課題のフォールバックを返します(description=%r)",
            clean,
            exc_info=True,
        )
        return {"spec": _fallback_spreadsheet_spec(lang), "_token_usage": no_usage}


def _sanitize_sheet_name(name: str, fallback: str) -> str:
    cleaned = _INVALID_SHEET_NAME_CHARS.sub("", (name or "").strip()) or fallback
    return cleaned[:31]


def build_xlsx_bytes(spec: dict) -> bytes:
    """generate_task_spreadsheetが返したJSON仕様から、実際の.xlsxバイト列を組み立てる。"""
    workbook = Workbook()
    workbook.remove(workbook.active)

    sheets = spec.get("sheets") or []
    for i, sheet_spec in enumerate(sheets):
        name = _sanitize_sheet_name(sheet_spec.get("name", ""), f"Sheet{i + 1}")
        worksheet = workbook.create_sheet(title=name)

        headers = sheet_spec.get("headers") or []
        if headers:
            worksheet.append([_sanitize_cell_value(v) for v in headers])

        for row in sheet_spec.get("rows") or []:
            worksheet.append([_sanitize_cell_value(v) for v in row])

    if not workbook.sheetnames:
        workbook.create_sheet(title="Sheet1")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
